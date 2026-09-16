"""
Genera el archivo plano de Requisiciones (INICIO + encabezados 440 + detalles 441 + FIN)
leyendo directamente el libro de Excel "Requisiciones.xlsm", SIN necesidad de abrir
Excel ni de macros/VBA.

Reproduce exactamente la misma regla que usa la macro "Generador de Archivos Planos"
que trae el libro:

    Fila 1 (desde columna B) -> nombre del campo
    Fila 2                    -> tipo: "Num" (numérico) o cualquier otra cosa (alfanumérico)
    Fila 3                    -> ancho del campo
    Fila 4 en adelante        -> datos (una fila de Excel = un registro del .txt)

        - Num -> se recorta y se rellena con CEROS a la izquierda
        - Alf -> se recorta y se rellena con ESPACIOS a la derecha

Como el generador lee la definición de campos directamente de las filas 1-3,
se adapta solo si mañana agregan columnas nuevas al layout (igual que pasó
entre la versión "02" y la "04" del formato de Requisiciones) - no hay que
tocar el código, solo agregar la columna en el Excel con su nombre/tipo/ancho.

IMPORTANTE - lo único que este script NO copia tal cual del Excel:

  1. El consecutivo global de línea (F_NUMERO-REG: las primeras 7 posiciones
     de TODO registro, incluidas las líneas de INICIO y FIN). Se recalcula
     automáticamente según el orden real en que se escribe el archivo.
     Se hizo así porque en el libro de ejemplo los consecutivos de la hoja
     DocInv (2, 3) no continúan con los de MovInv (1, 2, 3, 4, 5): si se
     copiaran tal cual, el archivo quedaría con la numeración de línea
     incorrecta y SIESA lo rechazaría. Todos los demás campos (tipo de
     documento, bodega, cantidades, referencias, etc.) se toman exactamente
     como estén escritos en el Excel.

  2. Las líneas de INICIO y FIN, que no vienen en ninguna hoja de este libro
     (en la herramienta original salían de hojas separadas "INICIO"/"FINAL"
     que no están en esta copia) - se arman con el mismo formato de 18
     caracteres que ya trae el archivo real: consecutivo(7) + tipo de
     registro "0000"/"9999" (4) + subtipo "00" (2) + versión "01" (2) +
     compañía (3).

CUIDADO con los campos de cantidad (f441_cant_base, f441_cant_2, etc.):
el layout espera 15 enteros + un PUNTO literal + 4 decimales
(ej. "000000000000078.0000"), pero la regla "Num" de este generador (igual
que la de la macro original) simplemente rellena con ceros a la izquierda
lo que haya en la celda - no inserta el punto por su cuenta. Si en la celda
solo se escribe "78", el resultado será "00000000000000000078" (SIN el
punto), lo cual el sistema destino interpretará mal. Por eso, en la celda
de cantidad hay que escribir ya el texto completo con el punto, por ejemplo:
    ="000000000000" & TEXTO(A1,"000") & ".0000"
o simplemente teclear el valor ya formateado como texto.

Requiere: pip install openpyxl
"""
from pathlib import Path

import openpyxl


def leer_definicion_campos(ws):
    """Lee nombre/tipo/ancho de cada campo (columna B en adelante) de una hoja."""
    campos = []
    col = 2  # la columna A es solo un marcador de fila, no se exporta
    while True:
        nombre = ws.cell(row=1, column=col).value
        if nombre is None or str(nombre).strip() == "":
            break
        tipo = ws.cell(row=2, column=col).value
        ancho = ws.cell(row=3, column=col).value
        campos.append({
            "nombre": str(nombre).strip(),
            "tipo": "Num" if str(tipo).strip().lower() == "num" else "Alf",
            "ancho": int(ancho),
        })
        col += 1
    return campos


def leer_filas_datos(ws, num_campos):
    """Lee los datos desde la fila 4 hasta la primera fila con la columna A vacía."""
    filas = []
    fila = 4
    while True:
        marcador = ws.cell(row=fila, column=1).value
        if marcador is None or str(marcador).strip() == "":
            break
        valores = [ws.cell(row=fila, column=col).value for col in range(2, 2 + num_campos)]
        filas.append(valores)
        fila += 1
    return filas


def formatear_valor(valor, tipo, ancho, ubicacion=""):
    """Aplica la regla Num=ceros a la izquierda / Alf=espacios a la derecha."""
    if valor is None:
        texto = ""
    elif isinstance(valor, float) and valor.is_integer():
        texto = str(int(valor))
    else:
        texto = str(valor).strip()

    if len(texto) > ancho:
        raise ValueError(
            f"{ubicacion}: el valor '{texto}' (largo {len(texto)}) "
            f"no cabe en el campo de {ancho} caracteres. Revisa el dato en el Excel."
        )

    return texto.zfill(ancho) if tipo == "Num" else texto.ljust(ancho)


def formatear_registro(valores, campos, ubicacion=""):
    return "".join(
        formatear_valor(v, c["tipo"], c["ancho"], f"{ubicacion} - campo '{c['nombre']}'")
        for v, c in zip(valores, campos)
    )


def generar_txt_de_hoja(ruta_xlsm, nombre_hoja, ruta_txt_salida):
    """
    Genera el .txt de UNA sola hoja, tal cual lo hacía la macro original para
    cualquiera de sus ~35 tipos de "plano" (TERCEROS, PROVEEDORES, ITEMS,
    CLIENTE, RUTAS, etc.): un archivo por hoja, sin combinar nada ni agregar
    INICIO/FIN. Sirve para cualquier hoja que tenga la misma convención
    (fila 1 = nombre, fila 2 = tipo, fila 3 = ancho, fila 4+ = datos).

    A diferencia de 'generar_txt_requisiciones', aquí SÍ se respeta tal cual
    el valor de cada celda (incluida la primera columna, ej. F_NUMERO-REG) -
    exactamente como hacía la macro. Si esa hoja necesita un consecutivo
    global correcto, hay que escribirlo bien en el propio Excel.
    """
    wb = openpyxl.load_workbook(ruta_xlsm, data_only=True)
    ws = wb[nombre_hoja]

    campos = leer_definicion_campos(ws)
    filas = leer_filas_datos(ws, len(campos))

    if not filas:
        raise ValueError(f"La hoja '{nombre_hoja}' no tiene datos desde la fila 4.")

    lineas = [
        formatear_registro(valores, campos, f"{nombre_hoja} fila {i}")
        for i, valores in enumerate(filas, start=4)
    ]

    contenido = "\n".join(lineas) + "\n"
    Path(ruta_txt_salida).write_text(contenido, encoding="latin-1")

    return ruta_txt_salida


def generar_txt_requisiciones(
    ruta_xlsm,
    ruta_txt_salida,
    hoja_encabezado="DocInv",
    hoja_detalle="MovInv",
    compania="001",
):
    """
    Arma el archivo completo (INICIO + encabezados + detalles + FIN) y lo
    escribe en 'ruta_txt_salida'. Retorna la ruta del archivo generado.
    """
    wb = openpyxl.load_workbook(ruta_xlsm, data_only=True)
    ws_enc = wb[hoja_encabezado]
    ws_det = wb[hoja_detalle]

    campos_enc = leer_definicion_campos(ws_enc)
    campos_det = leer_definicion_campos(ws_det)

    filas_enc = leer_filas_datos(ws_enc, len(campos_enc))
    filas_det = leer_filas_datos(ws_det, len(campos_det))

    if not filas_enc:
        raise ValueError(f"La hoja '{hoja_encabezado}' no tiene datos desde la fila 4.")
    if not filas_det:
        raise ValueError(f"La hoja '{hoja_detalle}' no tiene datos desde la fila 4.")

    lineas = []
    cons = 1
    lineas.append(f"{cons:07d}" + "0000" + "00" + "01" + compania.zfill(3))  # INICIO

    # Encabezados (440): el primer campo de cada hoja es F_NUMERO-REG, se
    # reemplaza por el consecutivo global; el resto se toma tal cual del Excel.
    for i, valores in enumerate(filas_enc, start=4):
        cons += 1
        resto = formatear_registro(valores[1:], campos_enc[1:], f"{hoja_encabezado} fila {i}")
        lineas.append(f"{cons:07d}{resto}")

    # Detalles (441)
    for i, valores in enumerate(filas_det, start=4):
        cons += 1
        resto = formatear_registro(valores[1:], campos_det[1:], f"{hoja_detalle} fila {i}")
        lineas.append(f"{cons:07d}{resto}")

    cons += 1
    lineas.append(f"{cons:07d}" + "9999" + "00" + "01" + compania.zfill(3))  # FIN

    contenido = "\n".join(lineas) + "\n"
    Path(ruta_txt_salida).write_text(contenido, encoding="latin-1")

    return ruta_txt_salida


if __name__ == "__main__":
    # Caso Requisiciones: un solo archivo combinado con INICIO + encabezados + detalles + FIN
    ruta = generar_txt_requisiciones(
        ruta_xlsm="Copy of Requisiciones.xlsm",
        ruta_txt_salida="REQUISICIONES_MASIVOS.txt",
    )
    print("Archivo generado en:", ruta)

    # Caso genérico: un archivo por hoja, igual que la macro original.
    # Sirve para cualquier otra hoja del libro (TERCEROS, PROVEEDORES, ITEMS, etc.)
    # generar_txt_de_hoja("Copy of Requisiciones.xlsm", "DocInv", "DocInv.TXT")
    # generar_txt_de_hoja("Copy of Requisiciones.xlsm", "MovInv", "MovInv.TXT")
